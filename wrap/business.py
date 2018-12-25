# -*- coding: utf-8 -*-

from wrap.base import *
from openpyxl import Workbook
from openpyxl import load_workbook
import datetime

class Business:

    def __init__(self):
        self.db = DB(db='wanbo')
        self._tables = {'测试':'test','收支明细':'balance','应收账款':'account','合同明细':'contract','开票明细':'invoice'}
        self._stats = ['应收账款统计','开票统计','收支统计','费用统计']
        self._statsAccount = {
            '今年合同额':{'row':0,'column':0,'form':'double'},
            '历年合同额':{'row':0,'column':1,'form':'double'},
            '合同总额':{'row':0,'column':2,'form':'double'},
            '今年回款额':{'row':2,'column':0,'form':'double'},
            '历年回款额':{'row':2,'column':1,'form':'double'},
            '回款总额':{'row':2,'column':2,'form':'double'},
            '今年未收款额':{'row':4,'column':0,'form':'double'},
            '历年未收款额':{'row':4,'column':1,'form':'double'},
            '未收款总额':{'row':4,'column':2,'form':'double'},
            '正常欠款':{'row':4,'column':3,'form':'double'},
            '异常欠款':{'row':4,'column':4,'form':'double'},
            '到期欠款':{'row':4,'column':5,'form':'double'},
            '今年坏账额':{'row':6,'column':0,'form':'double'},
            '历年坏账额':{'row':6,'column':1,'form':'double'},
            '坏账总额':{'row':6,'column':2,'form':'double'},
            '今年提成额':{'row':8,'column':0,'form':'double'},
            '历年提成额':{'row':8,'column':1,'form':'double'},
            '提成总额':{'row':8,'column':2,'form':'double'},
            '今年合同欠款率':{'row':10,'column':0,'form':'percent'},
            '历年合同欠款率':{'row':10,'column':1,'form':'percent'},
            '总欠款率':{'row':10,'column':2,'form':'percent'}
        }
        self._statsInvoice = (['月份','未税金额','税额','含税金额'],['str','double','double','double'])
        self._statsBalance = (['年份','来源','主营业务收入','其他业务收入','营业外收入','上年度余额转入','年度总收入','支出','结余','利润','年初余额','当前余额','余额增长','增长率'],
                              ['str','str','double','double','double','double','double','double','double','double','double','double','double','percent'])
        self._statsCost = None

    def loadExcel(self, fn):
        self.fn = fn
        self.wb = load_workbook(fn, read_only=True, data_only=True)
        return self.wb

    def closeExcel(self):
        self.wb.close()

    def commit(self):
        self.db.commit()

    def rollback(self):
        self.db.rollback()

    def __del__(self):
        if self.db != None:
            del self.db
            self.db = None

    def tables(self):
        return self._tables

    def stats(self):
        return self._stats

    def statsAccount(self):
        return self._statsAccount

    def statsInvoice(self):
        return self._statsInvoice

    def statsBalance(self):
        return self._statsBalance

    def statsCost(self):
        self._statsCost = ['一级类目','二级类目']
        sql = 'select date_format(date,"%Y-%m")as"月份" from balance where year(date)=year(now()) group by date_format(date,"%Y-%m")'
        tmp = self.db.query(sql)
        for t in tmp:
            self._statsCost.append(t['月份'])
        return self._statsCost

    def selectTableHead(self, table):
        sql = 'select %s_en,%s_zh,%s_tp from head where %s_en is not null' % (table,table,table,table)
        ret = self.db.query(sql)
        en = []
        zh = []
        tp = []
        for r in ret:
            en.append(r['%s_en' % table])
            zh.append(r['%s_zh' % table])
            tp.append(r['%s_tp' % table])
        return (en,zh,tp)

    def getInsertTemplates(self, table, enHead):
        prefix = 'insert into %s (' % table
        suffix = 'values('
        #编号由数据库自动生成，所以从1开始
        for n in range(1, len(enHead)):
            prefix += '%s,' % enHead[n]
            suffix += '%TBD,'
        return '%s) %s)' % (prefix.rstrip(','),suffix.rstrip(','))

    def selectTableData(self, table, condition=None):
        if condition == None:
            sql = 'select * from %s' % table
        else:
            sql = 'select * from %s where %s' % (table,condition)
        return self.db.query(sql)

    def updateTableById(self, table, field, tp, value, field_id, value_id):
        sql = 'update %s set %s = %TBD where %s = %d' 
        try:
            if value.strip() == '':
                sql = sql.replace('%TBD','%s', 1)
                value = 'NULL'
            elif tp == 'int':
                sql = sql.replace('%TBD','%d', 1)
                value = int(value)
            elif tp == 'double':
                sql = sql.replace('%TBD','%.2f', 1)
                value = float(value.replace(',',''))
            else:
                sql = sql.replace('%TBD','"%s"', 1)
        except:
            return False
        sql = sql % (table,field,value,field_id,value_id)
        return self.db.exec(sql)

    def insertTable(self, tableZh, itemData, commit=True):
        table = self.tables()[tableZh]
        head = self.selectTableHead(table)
        sql = self.getInsertTemplates(table, head[0])
        r = 0
        datas = []
        try:
            #编号由数据库自动生成，所以从1开始
            for c in range(1, len(itemData)):
                txt = itemData[c]
                if txt == 'NULL':
                    sql = sql.replace('%TBD','%s', 1)
                elif head[2][c] == 'int':
                    sql = sql.replace('%TBD','%d', 1)
                    txt = int(txt)
                elif head[2][c] == 'double':
                    sql = sql.replace('%TBD','%.2f', 1)
                    txt = float(txt.replace(',',''))
                else:
                    sql = sql.replace('%TBD','"%s"', 1)
                datas.append(txt)
        except:
            return False
        sql = sql % tuple(datas)
        return self.db.exec(sql, commit)

    def getContractAmount(self, alias, condition=None):
        if condition == None:
            sql = 'select sum(amount) as %s from contract' % alias
        else:
            sql = 'select sum(amount) as %s from contract where %s' % (alias,condition)
        return self.db.query(sql)

    def getInvoiceStats(self):
        sql = 'select date_format(date,"%Y-%m")as"月份",sum(price_notax)as"未税金额",sum(tax)as"税额" from invoice group by date_format(date,"%Y-%m")'
        ret = self.db.query(sql)
        tmp = {}
        tmp['月份'] = '合计'
        tmp['未税金额'] = 0.0
        tmp['税额'] = 0.0
        tmp['含税金额'] = 0.0
        i = 0
        for i in range(0,len(ret)):
            ret[i]['含税金额'] = round(ret[i]['未税金额']+ret[i]['税额'], 2)
            tmp['未税金额'] += ret[i]['未税金额']
            tmp['税额'] += ret[i]['税额']
            tmp['含税金额'] += ret[i]['含税金额']
        for k,v in tmp.items():
            if isinstance(v, float):
                tmp[k] = round(v, 2)
        ret.append(tmp)
        return ret

    def getBalanceStats(self):
        ret = []
        rowCount =  1
        total = {}
        to = {'主营业务收入':0.0,'其他业务收入':0.0,'营业外收入':0.0,'上年度余额转入':0.0,'年度总收入':0.0,'支出':0.0,'结余':0.0,'年初余额':0.0,'当前余额':0.0,'余额增长':0.0,'增长率':0.0}
        sql = 'select date_format(date,"%Y")as"year" from balance group by date_format(date,"%Y")'
        years = self.db.query(sql)
        #sql = 'select source from balance group by source'
        #sources = self.db.query(sql)
        sources = ['建设银行（基本户）','平安银行（姚洋）','平安银行（李昱平）']
        for y in years:
            rowCount += 1
            year = y['year']
            total[year] = to.copy()
            mp_year = {}
            lst_src = []
            for src in sources:
                mp = {}
                rowCount += 1
                #src = src['source']
                mp[src] = {}
                alias = '主营业务收入'
                sql = 'select sum(income)as"%s" from balance where date_format(date,"%%Y")="%s" and class2="%s" and source="%s"' % (alias,year,alias,src)
                tmp = self.db.query(sql)
                if tmp[0][alias] == None:
                    tmp[0][alias] = 0.0
                mp[src][alias] = tmp[0][alias]
                if src != '平安银行（李昱平）':
                    total[year][alias] += mp[src][alias]
                alias = '其他业务收入'
                sql = 'select sum(income)as"%s" from balance where date_format(date,"%%Y")="%s" and class2="%s" and source="%s"' % (alias,year,alias,src)
                tmp = self.db.query(sql)
                if tmp[0][alias] == None:
                    tmp[0][alias] = 0.0
                mp[src][alias] = tmp[0][alias]
                if src != '平安银行（李昱平）':
                    total[year][alias] += mp[src][alias]
                alias = '营业外收入'
                sql = 'select sum(income)as"%s" from balance where date_format(date,"%%Y")="%s" and class2="%s" and source="%s"' % (alias,year,alias,src)
                tmp = self.db.query(sql)
                if tmp[0][alias] == None:
                    tmp[0][alias] = 0.0
                mp[src][alias] = tmp[0][alias]
                if src != '平安银行（李昱平）':
                    total[year][alias] += mp[src][alias]
                alias = '上年度余额转入'
                sql = 'select sum(income)as"%s" from balance where date_format(date,"%%Y")="%s" and class2="%s" and source="%s"' % (alias,year,'余额转入',src)
                tmp = self.db.query(sql)
                if tmp[0][alias] == None:
                    tmp[0][alias] = 0.0
                mp[src][alias] = tmp[0][alias]
                total[year][alias] += mp[src][alias]
                total[year]['年初余额'] += mp[src][alias]
                alias = '年度总收入'
                sql = 'select sum(income)as"%s" from balance where date_format(date,"%%Y")="%s" and source="%s"' % (alias,year,src)
                tmp = self.db.query(sql)
                if tmp[0][alias] == None:
                    tmp[0][alias] = 0.0
                mp[src][alias] = round(tmp[0][alias]-mp[src]['上年度余额转入'], 2)
                if src != '平安银行（李昱平）':
                    total[year][alias] += mp[src][alias]
                alias = '支出'
                sql = 'select sum(pay)as"%s" from balance where date_format(date,"%%Y")="%s" and source="%s"' % (alias,year,src)
                tmp = self.db.query(sql)
                if tmp[0][alias] == None:
                    tmp[0][alias] = 0.0
                mp[src][alias] = tmp[0][alias]
                if src != '平安银行（李昱平）':
                    total[year][alias] += mp[src][alias]
                alias = '结余'
                sql = 'select sum(income)-sum(pay)as"%s" from balance where date_format(date,"%%Y")="%s" and source="%s"' % (alias,year,src)
                tmp = self.db.query(sql)
                if tmp[0][alias] == None:
                    tmp[0][alias] = 0.0
                mp[src][alias] = tmp[0][alias]
                total[year][alias] += mp[src][alias]
                total[year]['当前余额'] += mp[src][alias]
                lst_src.append(mp)
            total[year]['余额增长'] = total[year]['当前余额'] - total[year]['年初余额']
            total[year]['增长率'] = total[year]['余额增长'] / total[year]['年初余额']
            total[year]['利润'] = total[year]['年度总收入'] - total[year]['支出']
            mp_year[year] = lst_src
            ret.append(mp_year)
        for v1 in total.values():
            for k2,v2 in v1.items():
                v1[k2] = round(v2, 2)
        return (ret, rowCount, total)

    def getCostStats(self):
        mp = {}
        #总费用
        sql = 'select sum(pay)as"费用",date_format(date,"%Y-%m")as"月份" from balance where year(date)=year(now()) group by date_format(date,"%Y-%m")'
        tmp = self.db.query(sql)
        mp['总费用汇总'] = {}
        mp['总费用汇总']['费用'] = tmp
        count = 1   #统计一级类目和二级类目的总数
        #一级类目
        sql = 'select class1 from balance where year(date)=year(now()) group by class1'
        tmp1 = self.db.query(sql)
        for t1 in tmp1:
            count += 1
            class1 = t1['class1']
            mp[class1] = {}
            #一级类目的费用
            sql = 'select date_format(date,"%%Y-%%m")as"月份",sum(pay)as"费用" from balance where year(date)=year(now()) and class1="%s" group by date_format(date,"%%Y-%%m")' % class1
            tmp2 = self.db.query(sql)
            mp[class1]['费用'] = []
            for t2 in tmp2:
                mp[class1]['费用'].append(t2)
            mp[class1]['二级类目'] = {}
            #一级类目下的二级类目
            sql = 'select class2 from balance where year(date)=year(now()) and class1="%s" group by class2' % class1
            tmp3 = self.db.query(sql)
            for t3 in tmp3:
                count += 1
                class2 = t3['class2']
                mp[class1]['二级类目'][class2] = []
                #二级类目的费用
                sql = 'select date_format(date,"%%Y-%%m")as"月份",sum(pay)as"费用" from balance where year(date)=year(now()) and class1="%s" and class2="%s" group by date_format(date,"%%Y-%%m")'\
                        % (class1,class2)
                tmp4 = self.db.query(sql)
                for t4 in tmp4:
                    mp[class1]['二级类目'][class2].append(t4)
        return (mp,count)

    def getAccountStats(self):
        mp = {}
        sql = 'select sum(amount)as"合同总额",sum(paid)as"回款总额",sum(unpaid)as"未收款总额",sum(debt)as"坏账总额",sum(commission)as"提成总额" from account'
        mp.update(self.db.query(sql)[0])
        sql = 'select sum(amount)as"今年合同额",sum(paid)as"今年回款额",sum(unpaid)as"今年未收款额",sum(debt)as"今年坏账额",sum(commission)as"今年提成额" from account where year(date)=year(now())'
        mp.update(self.db.query(sql)[0])
        sql = 'select sum(amount)as"历年合同额",sum(paid)as"历年回款额",sum(unpaid)as"历年未收款额",sum(debt)as"历年坏账额",sum(commission)as"历年提成额" from account where year(date)<year(now())'
        mp.update(self.db.query(sql)[0])
        sql = 'select sum(unpaid)as"正常欠款" from account where status="正常欠款"'
        mp.update(self.db.query(sql)[0])
        sql = 'select sum(unpaid)as"异常欠款" from account where status="异常欠款"'
        mp.update(self.db.query(sql)[0])
        sql = 'select sum(unpaid)as"到期欠款" from account where status="到期欠款"'
        mp.update(self.db.query(sql)[0])
        for k,v in mp.items():
            if v == None:
                mp[k] = 0.0
        return mp

    def TestDB(self):
        sql = 'select * from test'
        print(self.db.query(sql))
    
    def insertToTable(self, mp, table, commit=True):
        keys_sql = '('
        values_sql = '('
        for k,v in mp.items():
            keys_sql += '%s,' % k
            if v==None or v=='-' or v=='/' or v=='?' or v=='？':
                values_sql += 'NULL,'
            elif isinstance(v, str):
                if v.strip() == '':
                    values_sql += 'NULL,'
                else:
                    values_sql += '"%s",' % v
            elif isinstance(v, datetime.datetime):
                values_sql += '"%s",' % v.strftime('%Y-%m-%d')
            else:
                values_sql += '%s,' % v
        keys_sql = keys_sql.rstrip(',')
        values_sql = values_sql.rstrip(',')
        keys_sql += ')'
        values_sql += ')'
        sql = 'insert into %s %s values%s' % (table,keys_sql,values_sql)
        return self.db.exec(sql, commit)
    
    def ReadBalanceData(self, ws, source):
        n = 0
        mp = {}
        self.db.resetCount()
        ok = True
        try:
            for row in ws.rows:
                n = n + 1
                if n <= 3:
                    continue
                if len(row) < 10:
                    continue
                mp['class1'] = row[0].value
                if mp['class1']==None or mp['class1'].strip()=='':
                    continue
                mp['class2'] = row[1].value
                year = row[2].value
                month = row[3].value
                day = row[4].value
                mp['date'] = '%d-%d-%d' % (year,month,day)
                mp['abstract'] = row[5].value
                mp['income'] = row[6].value
                mp['pay'] = row[7].value
                mp['balance'] = row[8].value
                mp['type'] = row[9].value
                mp['credence'] = row[10].value
                mp['remark'] = row[11].value
                mp['source'] = source
                if self.insertToTable(mp, 'balance', False) == False:
                    ok = False
                    break
        except:
            ok = False
        if ok == False:
            self.db.rollback()
            GL.setErr('导入 %s 第 %d 行时报错！' % (ws.title,n))
            return False
        else:
            self.db.commit()
            return self.db.getCount()
    
    def ReadAccountData(self, ws):
        n = 0
        mp = {}
        self.db.resetCount()
        ok = True
        try:
            for row in ws.rows:
                n = n + 1
                if n <= 4:
                    continue
                if len(row) < 18:
                    continue
                mp['no'] = row[0].value
                if mp['no']==None or isinstance(mp['no'],int)==False:
                    continue
                del mp['no']
                year = row[1].value
                month = row[2].value
                day = row[3].value
                mp['date'] = '%d-%d-%d' % (year,month,day)
                mp['seller'] = row[4].value
                mp['contract'] = row[5].value
                mp['client'] = row[6].value
                mp['amount'] = row[7].value
                mp['paid'] = row[8].value
                mp['unpaid'] = row[9].value
                mp['debt'] = row[10].value
                mp['percent'] = row[11].value
                mp['invoice_type'] = row[12].value
                mp['invoice_paid'] = row[13].value
                mp['invoice_unpaid'] = row[14].value
                mp['status'] = row[15].value
                mp['unpaid_reason'] = row[16].value
                mp['commission'] = row[17].value
                mp['commission_date'] = row[18].value
                mp['remark'] = row[19].value
                mp['invoice_commission'] = row[20].value
                if self.insertToTable(mp, 'account', False) == False:
                    ok = False
                    break
        except:
            ok = False
        if ok == False:
            self.db.rollback()
            GL.setErr('导入 %s 第 %d 行时报错！' % (ws.title,n))
            return False
        else:
            self.db.commit()
            return self.db.getCount()
    
    def ReadContractData(self, ws):
        n = 0
        mp = {}
        self.db.resetCount()
        ok = True
        try:
            for row in ws.rows:
                n = n + 1
                if n <= 4:
                    continue
                if len(row) < 18:
                    continue
                mp['contract'] = row[1].value
                if mp['contract']==None or mp['contract'].strip()=='':
                    continue
                mp['date'] = row[2].value
                mp['seller'] = row[3].value
                mp['client'] = row[4].value
                mp['product'] = row[5].value
                mp['model'] = row[6].value
                mp['unit'] = row[7].value
                mp['quantity'] = row[8].value
                mp['price'] = row[9].value
                mp['count'] = row[10].value
                mp['amount'] = row[11].value
                mp['tax_rate'] = row[12].value
                mp['date_plan'] = row[13].value
                mp['date_actual'] = row[14].value
                mp['ontime'] = row[15].value
                mp['ontime_whynot'] = row[16].value
                mp['project'] = row[17].value
                mp['province'] = row[18].value
                mp['remark'] = row[19].value
                if self.insertToTable(mp, 'contract', False) == False:
                    ok = False
                    break
        except:
            ok = False
        if ok == False:
            self.db.rollback()
            GL.setErr('导入 %s 第 %d 行时报错！' % (ws.title,n))
            return False
        else:
            self.db.commit()
            return self.db.getCount()
    
    def ReadInvoiceData(self, ws):
        n = 0
        mp = {}
        self.db.resetCount()
        ok = True
        try:
            for row in ws.rows:
                n = n + 1
                if n <= 2:
                    continue
                if len(row) < 16:
                    continue
                mp['contract'] = row[0].value
                if mp['contract']==None or mp['contract'].strip()=='':
                    continue
                year = row[1].value
                month = row[2].value
                day = row[3].value
                mp['date'] = '%d-%d-%d' % (year,month,day)
                mp['invoice_no'] = row[4].value
                mp['invoice_client'] = row[5].value
                mp['invoice_content'] = row[6].value
                mp['model'] = row[7].value
                mp['quantity'] = row[8].value
                mp['price'] = row[9].value
                mp['price_notax'] = row[10].value
                mp['tax'] = row[11].value
                mp['amount'] = row[12].value
                mp['invoice_amount'] = row[13].value
                mp['invoice_type'] = row[14].value
                mp['status'] = row[15].value
                mp['refund'] = row[16].value
                mp['remark'] = row[17].value
                if self.insertToTable(mp, 'invoice', False) == False:
                    ok = False
                    break
        except:
            ok = False
        if ok == False:
            self.db.rollback()
            GL.setErr('导入 %s 第 %d 行时报错！' % (ws.title,n))
            return False
        else:
            self.db.commit()
            return self.db.getCount()

